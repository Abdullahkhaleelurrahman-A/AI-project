import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def parse_markdown_to_doc(doc, md_content):
    lines = md_content.split('\n')
    in_code_block = False
    in_table = False
    table_rows = []
    
    for line in lines:
        line_stripped = line.strip()
        
        # Code blocks
        if line_stripped.startswith('```'):
            in_code_block = not in_code_block
            continue
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style = 'No Spacing'
            if p.runs:
                p.runs[0].font.name = 'Courier New'
            continue
            
        # Tables
        if line_stripped.startswith('|'):
            in_table = True
            if '---' not in line_stripped:
                cols = [c.strip() for c in line_stripped.split('|')[1:-1]]
                table_rows.append(cols)
            continue
        elif in_table:
            # We exited the table
            if len(table_rows) > 0:
                table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                table.style = 'Table Grid'
                for r_idx, row_data in enumerate(table_rows):
                    for c_idx, cell_data in enumerate(row_data):
                        if c_idx < len(table.columns):
                            # Remove bold markdown for simplicity in tables
                            clean_cell = cell_data.replace('**', '')
                            table.cell(r_idx, c_idx).text = clean_cell
                doc.add_paragraph() # spacing
            in_table = False
            table_rows = []
            if not line_stripped: continue
            
        if not line_stripped:
            continue
            
        # Headings
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)
            
        # Lists
        elif line.startswith('- ') or line.startswith('* '):
            text = line[2:]
            doc.add_paragraph(text, style='List Bullet')
        elif re.match(r'^\d+\.\s', line):
            text = line.split('. ', 1)[1]
            doc.add_paragraph(text, style='List Number')
            
        # Blockquotes/Alerts
        elif line.startswith('> '):
            p = doc.add_paragraph(line[2:])
            p.style = 'Quote'
            
        # Normal paragraph (with bold parsing)
        else:
            p = doc.add_paragraph()
            # Simple bold parsing
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    # Strip inline code ticks
                    part = part.replace('`', '')
                    p.add_run(part)

def create_word_document():
    doc = Document()
    
    # --- Title Page ---
    for _ in range(5):
        doc.add_paragraph()
    
    title = doc.add_paragraph("MARUDAM")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if title.runs:
        title.runs[0].font.size = Pt(48)
        title.runs[0].bold = True
    else:
        run = title.add_run("MARUDAM")
        run.font.size = Pt(48)
        run.bold = True
    
    subtitle = doc.add_paragraph("AI-Powered Crop Monitoring & Field Intelligence Platform")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle.runs:
        subtitle.runs[0].font.size = Pt(18)
    else:
        run = subtitle.add_run("AI-Powered Crop Monitoring & Field Intelligence Platform")
        run.font.size = Pt(18)
        
    doc.add_paragraph()
    
    doc_type = doc.add_paragraph("COMPLETE WEBSITE & SYSTEM DOCUMENTATION")
    doc_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if doc_type.runs:
        doc_type.runs[0].font.size = Pt(14)
        doc_type.runs[0].bold = True
    else:
        run = doc_type.add_run("COMPLETE WEBSITE & SYSTEM DOCUMENTATION")
        run.font.size = Pt(14)
        run.bold = True
        
    doc.add_paragraph()
    
    info = doc.add_paragraph("Current Implementation Documentation\nTechnology Stack: Next.js, FastAPI, Supabase, ESP32, Gemini AI\nGenerated: September 2026")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()
    
    # --- Executive Summary ---
    doc.add_heading("EXECUTIVE SUMMARY", level=1)
    doc.add_paragraph("MARUDAM is an advanced agricultural intelligence platform that bridges the gap between raw hardware telemetry and actionable agronomic insights. By combining ESP32-based hardware arrays, a FastAPI Python edge-processing layer, and a Next.js frontend, MARUDAM monitors real-time field conditions. Instead of relying on static, generic thresholds which fail across diverse microclimates, MARUDAM implements an adaptive baseline engine to learn the unique normal state of a specific field. When anomalies are detected, they are tracked as stateful events by a deterministic risk engine, and finally analyzed by Google Gemini AI, which cross-references the anomalies with hyper-localized Indian geographic data and dataset-injected pH/EC references to provide highly specific farming recommendations.")
    doc.add_page_break()
    
    # --- Mapping Markdown Files to Word Sections ---
    files_to_include = [
        "ARCHITECTURE.md",
        "WEBSITE_FLOWS.md",
        "FRONTEND_FUNCTIONS.md",
        "BACKEND_FUNCTIONS.md",
        "ESP32_AND_SENSORS.md",
        "SENSOR_DATA_PIPELINE.md",
        "DATABASE.md",
        "API_REFERENCE.md",
        "ADAPTIVE_BASELINE.md",
        "FORMULAS_AND_CALCULATIONS.md",
        "ANOMALY_DETECTION.md",
        "EVENTS.md",
        "RISK_ENGINE.md",
        "RECOMMENDATIONS.md",
        "AI_GEMINI.md",
        "LOCALIZATION.md",
        "SECURITY.md",
        "IMPLEMENTATION_STATUS.md",
        "KNOWN_ISSUES.md"
    ]
    
    for filename in files_to_include:
        filepath = os.path.join("docs", filename)
        if os.path.exists(filepath):
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()
                parse_markdown_to_doc(doc, content)
                doc.add_page_break()
                
    doc.save("docs/MARUDAM_COMPLETE_WEBSITE_DOCUMENTATION.docx")
    print("Successfully created docs/MARUDAM_COMPLETE_WEBSITE_DOCUMENTATION.docx")

if __name__ == "__main__":
    create_word_document()
